"""Autentica en Consuman y extrae Ordenes de Trabajo a JSON.

Esta etapa conserva el flujo de login normal. La llamada a la API se ejecuta con
``fetch`` dentro de la pagina autenticada, por lo que usa automaticamente las
cookies de la sesion creada por Playwright.
"""

from __future__ import annotations

import json
import os
import re
import sys
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from typing import Any

from order_status import status_name
from task_details import build_task_records

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from playwright.sync_api import (
    Error as PlaywrightError,
    TimeoutError as PlaywrightTimeoutError,
    sync_playwright,
)


LOGIN_URL = "https://panoi.consumannet.com.ar/Login.aspx"
WORK_ORDERS_URL = "https://panoi.consumannet.com.ar/Operational/WorkOrders.aspx"
API_URL = "https://panoi.consumannet.com.ar/api/workOrder/get-grid"
TASK_API_URL = "https://panoi.consumannet.com.ar/api/workOrderTask/get-grid"
OUTPUT_PATH = Path(__file__).resolve().parent / "ordenes_trabajo.json"
TASK_OUTPUT_PATH = Path(__file__).resolve().parent / "tareas_trabajo.json"
EXCEL_PATH = Path(__file__).resolve().parent / "ordenes_trabajo_Uruguay.xlsx"
TIMEOUT_MS = 60_000
LOGIN_TIMEOUT_MS = 90_000
LOGIN_ATTEMPTS = 3


class ConsumanError(RuntimeError):
    """Error legible durante la automatizacion de Consuman."""


def read_credentials() -> tuple[str, str]:
    username = os.getenv("CONSUMAN_USUARIO", "").strip()
    password = os.getenv("CONSUMAN_PASSWORD", "")
    missing = [
        name
        for name, value in (
            ("CONSUMAN_USUARIO", username),
            ("CONSUMAN_PASSWORD", password),
        )
        if not value
    ]
    if missing:
        raise ConsumanError(
            "Faltan variables de entorno: "
            + ", ".join(missing)
            + ". Definalas antes de ejecutar."
        )
    return username, password


def read_date(name: str, default: date) -> date:
    raw = os.getenv(name, "").strip()
    if not raw:
        return default
    try:
        return date.fromisoformat(raw)
    except ValueError as exc:
        raise ConsumanError(f"{name} debe usar el formato AAAA-MM-DD: {raw!r}") from exc


def iso_utc(day: date, end_of_day: bool = False) -> str:
    clock = time(23, 59, 59, 999000) if end_of_day else time.min
    value = datetime.combine(day, clock, tzinfo=timezone.utc)
    return value.isoformat(timespec="milliseconds").replace("+00:00", "Z")


def iso_datetime_utc(value: datetime) -> str:
    return value.astimezone(timezone.utc).isoformat(timespec="milliseconds").replace(
        "+00:00", "Z"
    )


def build_payload() -> dict[str, Any]:
    now = datetime.now().astimezone()
    today = now.date()
    from_date = read_date("CONSUMAN_DESDE", today - timedelta(days=30))
    to_date = read_date("CONSUMAN_HASTA", today + timedelta(days=30))
    if from_date > to_date:
        raise ConsumanError("CONSUMAN_DESDE no puede ser posterior a CONSUMAN_HASTA.")

    if os.getenv("CONSUMAN_DESDE", "").strip():
        from_value = datetime.combine(from_date, time.min).astimezone()
    else:
        from_value = now - timedelta(days=30)
    if os.getenv("CONSUMAN_HASTA", "").strip():
        to_value = datetime.combine(to_date, time(23, 59, 59, 999000)).astimezone()
    else:
        to_value = now + timedelta(days=30)

    return {
        "fromDate": iso_datetime_utc(from_value),
        "toDate": iso_datetime_utc(to_value),
        "daysAfter": 30,
        "daysBefore": 30,
        "equipment": [],
        "project": [],
        "staff": [],
        "taskType": [],
        "periodRange": True,
        # Incluye todos los estados de OT para no omitir ordenes recientes,
        # canceladas o en transiciones que la vista predeterminada no marca.
        "woState": [{"id": value} for value in (1, 2, 3, 4, 5)],
        "woStatesTask": [{"id": value} for value in (0, 2, 3)],
        "period": [{"id": 6}],
        "layout": [],
        "loadOptions": {
            "skip": 0,
            "take": 25,
            "requireTotalCount": True,
            "sort": [
                {"selector": "date", "desc": True},
                {"selector": "type", "desc": True},
            ],
        },
    }


def is_login_page(page) -> bool:
    return (
        "login.aspx" in page.url.lower()
        or page.locator('input[name="username"]').count() > 0
    )


def login(page, username: str, password: str) -> None:
    last_detail = "Consuman no completo el inicio de sesion."
    for attempt in range(1, LOGIN_ATTEMPTS + 1):
        print(f"Abriendo login: {LOGIN_URL} (intento {attempt}/{LOGIN_ATTEMPTS})")
        try:
            page.goto(
                LOGIN_URL,
                wait_until="domcontentloaded",
                timeout=LOGIN_TIMEOUT_MS,
            )
            user_input = page.locator('input[name="username"]')
            password_input = page.locator('input[name="password"]')
            login_button = page.locator("button", has_text="Login")
            user_input.wait_for(state="visible", timeout=LOGIN_TIMEOUT_MS)
            password_input.wait_for(state="visible", timeout=LOGIN_TIMEOUT_MS)
            login_button.wait_for(state="visible", timeout=LOGIN_TIMEOUT_MS)
            user_input.fill(username)
            password_input.fill(password)
            login_button.click()
            user_input.wait_for(state="hidden", timeout=LOGIN_TIMEOUT_MS)
            page.wait_for_load_state("domcontentloaded", timeout=LOGIN_TIMEOUT_MS)
            if not is_login_page(page):
                return

            messages = page.locator(
                ".alert-danger:visible, .validation-summary-errors:visible, "
                ".swal2-html-container:visible"
            ).all_inner_texts()
            if messages:
                raise ConsumanError(
                    "Consuman rechazo el acceso: " + " ".join(messages).strip()
                )
            last_detail = "Consuman permanecio en la pantalla de ingreso."
        except ConsumanError:
            raise
        except (PlaywrightTimeoutError, PlaywrightError) as exc:
            last_detail = f"{type(exc).__name__}: {str(exc).splitlines()[0]}"

        if attempt < LOGIN_ATTEMPTS:
            delay_ms = attempt * 5_000
            print(
                f"ADVERTENCIA: fallo transitorio de acceso. "
                f"Reintentando en {delay_ms // 1000} segundos..."
            )
            try:
                page.goto("about:blank", wait_until="domcontentloaded", timeout=10_000)
            except PlaywrightError:
                pass
            page.wait_for_timeout(delay_ms)

    raise ConsumanError(
        f"No se pudo iniciar sesion en Consuman despues de {LOGIN_ATTEMPTS} "
        f"intentos. Ultimo detalle: {last_detail}"
    )


def extract_orders(page, payload: dict[str, Any]) -> Any:
    observed_requests: list[dict[str, Any]] = []

    def observe_request(request) -> None:
        if "/api/workorder/get-grid" not in request.url.lower():
            return
        safe_headers = {
            key: value
            for key, value in request.headers.items()
            if key.lower() in {
                "accept",
                "content-type",
                "requestverificationtoken",
                "x-requested-with",
                "x-xsrf-token",
            }
        }
        observed_requests.append(
            {
                "method": request.method,
                "headers": safe_headers,
                "postData": request.post_data,
            }
        )

    page.on("request", observe_request)
    print(f"Navegando a Ordenes de Trabajo: {WORK_ORDERS_URL}")
    response = page.goto(WORK_ORDERS_URL, wait_until="domcontentloaded")
    page.locator("body").wait_for(state="visible")
    if response is not None and response.status >= 400:
        raise ConsumanError(f"La pagina devolvio HTTP {response.status}.")
    if is_login_page(page):
        raise ConsumanError("La sesion expiro o no quedo autenticada.")

    # Da tiempo a que la aplicacion haga su propia carga inicial. La observacion
    # solo conserva encabezados no sensibles y nunca copia cookies.
    page.wait_for_timeout(3_000)

    print(f"Consultando API para {payload['fromDate']} a {payload['toDate']}...")
    result = page.evaluate(
        """async ({ apiUrl, payload }) => {
            const allRows = [];
            let totalCount = null;
            let skip = 0;
            const take = 25;
            while (totalCount === null || skip < totalCount) {
                const pagePayload = {
                    ...payload,
                    loadOptions: { ...payload.loadOptions, skip, take, requireTotalCount: true },
                };
                const response = await fetch(apiUrl, {
                    method: "POST",
                    credentials: "same-origin",
                    headers: {
                        "Accept": "application/json, text/javascript, */*; q=0.01",
                        "Content-Type": "application/json",
                        "X-Requested-With": "XMLHttpRequest",
                    },
                    body: JSON.stringify(pagePayload),
                });
                const bodyText = await response.text();
                let data = null;
                try { data = JSON.parse(bodyText); } catch (_) {}
                if (!response.ok || data === null) {
                    return {
                        ok: false,
                        status: response.status,
                        statusText: response.statusText,
                        contentType: response.headers.get("content-type"),
                        data,
                        bodyPreview: bodyText.slice(0, 500),
                        failedSkip: skip,
                    };
                }
                const rows = Array.isArray(data) ? data : (data.data || data.rows || []);
                if (!Array.isArray(rows)) {
                    return { ok: false, status: 200, statusText: "Formato inesperado", data: null, bodyPreview: bodyText.slice(0, 500), failedSkip: skip };
                }
                allRows.push(...rows);
                totalCount = Number.isFinite(Number(data.totalCount)) ? Number(data.totalCount) : allRows.length;
                if (rows.length === 0 || rows.length < take) break;
                skip += take;
            }
            return {
                ok: true,
                status: 200,
                statusText: "OK",
                contentType: "application/json",
                data: { data: allRows, totalCount: totalCount ?? allRows.length },
                bodyPreview: "",
            };
        }""",
        {"apiUrl": API_URL, "payload": payload},
    )
    if not result["ok"]:
        diagnostic_path = OUTPUT_PATH.with_name("diagnostico_get_grid.json")
        diagnostic = {
            "attemptedPayload": payload,
            "httpStatus": result["status"],
            "bodyPreview": result["bodyPreview"],
            "observedPageRequests": observed_requests,
        }
        with diagnostic_path.open("w", encoding="utf-8", newline="\n") as file:
            json.dump(diagnostic, file, ensure_ascii=False, indent=2)
            file.write("\n")
        raise ConsumanError(
            f"La API devolvio HTTP {result['status']} {result['statusText']}. "
            f"Respuesta: {result['bodyPreview']!r}. Se guardo un diagnostico "
            f"sin cookies en {diagnostic_path}."
        )
    if result["data"] is None:
        raise ConsumanError(
            "La API no devolvio JSON. Content-Type: "
            + str(result.get("contentType"))
        )
    return result["data"]


def extract_tasks(page, payload: dict[str, Any]) -> Any:
    """Obtiene las tareas reales de la grilla Detalles de Consuman."""
    task_payload = {
        **payload,
        # La exportacion debe incluir tareas realizadas, pendientes, canceladas
        # y sin realizar. La vista predeterminada omite las realizadas.
        "woStatesTask": [{"id": value} for value in (0, 1, 2, 3)],
        "loadOptions": {
            "skip": 0,
            "take": 100,
            "requireTotalCount": True,
            "sort": [{"selector": "plannedDate", "desc": True}],
        },
    }
    print("Consultando tareas reales de la grilla Detalles...")
    result = page.evaluate(
        """async ({ apiUrl, payload }) => {
            const allRows = [];
            let totalCount = null;
            let skip = 0;
            const take = 100;
            while (totalCount === null || skip < totalCount) {
                const pagePayload = {
                    ...payload,
                    loadOptions: { ...payload.loadOptions, skip, take, requireTotalCount: true },
                };
                const response = await fetch(apiUrl, {
                    method: "POST",
                    credentials: "same-origin",
                    headers: {
                        "Accept": "application/json, text/javascript, */*; q=0.01",
                        "Content-Type": "application/json",
                        "X-Requested-With": "XMLHttpRequest",
                    },
                    body: JSON.stringify(pagePayload),
                });
                const bodyText = await response.text();
                let data = null;
                try { data = JSON.parse(bodyText); } catch (_) {}
                if (!response.ok || data === null) {
                    return { ok: false, status: response.status, bodyPreview: bodyText.slice(0, 500) };
                }
                const rows = Array.isArray(data) ? data : (data.data || data.rows || []);
                if (!Array.isArray(rows)) {
                    return { ok: false, status: 200, bodyPreview: bodyText.slice(0, 500) };
                }
                allRows.push(...rows);
                totalCount = Number.isFinite(Number(data.totalCount)) ? Number(data.totalCount) : allRows.length;
                if (rows.length === 0 || rows.length < take) break;
                skip += take;
            }
            return { ok: true, data: { data: allRows, totalCount: totalCount ?? allRows.length } };
        }""",
        {"apiUrl": TASK_API_URL, "payload": task_payload},
    )
    if not result["ok"]:
        raise ConsumanError(
            f"No se pudieron obtener las tareas (HTTP {result['status']}): "
            f"{result['bodyPreview']!r}"
        )
    return result["data"]


def task_columns(page) -> list[dict[str, Any]]:
    """Lee el campo real que usa Fecha realización en la grilla Detalles."""
    page.get_by_text(re.compile(r"^\s*detalles\s*$", re.IGNORECASE)).click()
    page.wait_for_function("""() => {
        const clean = x => String(x || '').normalize('NFD').replace(/[\\u0300-\\u036f]/g, '').toLowerCase();
        return [...document.querySelectorAll('.dx-header-row td')]
          .some(cell => clean(cell.textContent).includes('fecha realizacion'));
    }""", timeout=TIMEOUT_MS)
    columns = page.evaluate("""() => {
        const found = [];
        const flatten = cols => (cols || []).forEach(col => {
            if (col && typeof col === 'object') {
                if (col.caption && col.dataField) found.push({caption: col.caption, dataField: col.dataField});
                flatten(col.columns);
            }
        });
        for (const element of document.querySelectorAll('.dx-datagrid, .dx-treelist, .dx-widget')) {
            if (!element.getClientRects().length) continue;
            // En DevExtreme 19 la clase puede estar en el contenido del widget.
            for (const host of [element, element.parentElement, element.parentElement && element.parentElement.parentElement]) {
                if (!host) continue;
                let grid = null;
                try {
                    if (window.DevExpress && window.DevExpress.ui && window.DevExpress.ui.dxDataGrid)
                        grid = window.DevExpress.ui.dxDataGrid.getInstance(host);
                    if (!grid && window.jQuery && window.jQuery.fn.dxDataGrid)
                        grid = window.jQuery(host).dxDataGrid('instance');
                } catch (_) { continue; }
                if (grid) {
                    // getVisibleColumns contiene los captions ya resueltos por la aplicación.
                    flatten(grid.getVisibleColumns());
                    flatten(grid.option('columns'));
                    break;
                }
            }
        }
        return found;
    }""")
    if not columns:
        print("Diagnóstico de grillas:", page.evaluate("""() => [...document.querySelectorAll('.dx-datagrid, .dx-widget')].filter(e => e.getClientRects().length).slice(0,12).map(e => ({id:e.id,classes:e.className,parent:e.parentElement && e.parentElement.id}))"""))
        raise ConsumanError("No se pudo leer la configuración de la grilla Detalles.")
    return columns


def add_task_names(orders_data: Any, tasks_data: Any) -> Any:
    orders = orders_data.get("data", []) if isinstance(orders_data, dict) else orders_data
    tasks = tasks_data.get("data", []) if isinstance(tasks_data, dict) else tasks_data
    if not isinstance(orders, list) or not isinstance(tasks, list):
        raise ConsumanError("Consuman devolvio un formato inesperado al combinar tareas.")

    tasks_by_order: dict[str, list[str]] = {}
    task_states_by_order: dict[str, list[int]] = {}
    for task in tasks:
        if not isinstance(task, dict):
            continue
        key = str(task.get("workOrderId") or task.get("workOrderNumber") or "")
        name = str(task.get("task") or "").strip()
        if not key or not name:
            continue
        names = tasks_by_order.setdefault(key, [])
        if name not in names:
            names.append(name)
        raw_state = task.get("state")
        if raw_state is not None:
            try:
                state = int(raw_state)
            except (TypeError, ValueError):
                state = -1
            if state >= 0:
                states = task_states_by_order.setdefault(key, [])
                if state not in states:
                    states.append(state)

    for order in orders:
        if not isinstance(order, dict):
            continue
        key = str(order.get("id") or order.get("number") or "")
        names = tasks_by_order.get(key, [])
        order["task"] = " | ".join(names)
        order["tasks"] = names
        order["taskStates"] = task_states_by_order.get(key, [])
        order["sheetStatus"] = status_name(order)
    return orders_data


def count_records(data: Any) -> int | None:
    if isinstance(data, list):
        return len(data)
    if isinstance(data, dict):
        for key in ("data", "rows", "items", "result", "workOrders"):
            if isinstance(data.get(key), list):
                return len(data[key])
    return None


def save_json_atomic(data: Any) -> None:
    temp_path = OUTPUT_PATH.with_suffix(".json.tmp")
    with temp_path.open("w", encoding="utf-8", newline="\n") as file:
        json.dump(data, file, ensure_ascii=False, indent=2)
        file.write("\n")
    os.replace(temp_path, OUTPUT_PATH)


def task_name(order: dict[str, Any]) -> str:
    tasks = order.get("tasks")
    if isinstance(tasks, list):
        names = [str(value).strip() for value in tasks if str(value).strip()]
        if names:
            return " | ".join(dict.fromkeys(names))
    return str(order.get("task") or "").strip()


def excel_date(value: Any) -> datetime | str | None:
    if not value:
        return None
    text = str(value).strip()
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
        if parsed.tzinfo is not None:
            parsed = parsed.astimezone().replace(tzinfo=None)
        return parsed
    except ValueError:
        return text


def generate_excel(data: Any) -> None:
    orders = data.get("data", []) if isinstance(data, dict) else data
    if not isinstance(orders, list):
        raise ConsumanError("No se encontro una lista de ordenes para generar el Excel.")

    temp_excel = EXCEL_PATH.with_name("ordenes_trabajo_Uruguay.nuevo.xlsx")
    try:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Ordenes de trabajo"
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"

        headers = ["OT", "Activo", "Tarea", "Fecha planificada", "Estado"]
        sheet.append(headers)
        for order in orders:
            if not isinstance(order, dict):
                continue
            sheet.append(
                [
                    str(order.get("number") or ""),
                    str(order.get("equipment") or ""),
                    task_name(order),
                    excel_date(order.get("date")),
                    str(order.get("sheetStatus") or status_name(order)),
                ]
            )

        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(color="FFFFFF", bold=True)
        header_border = Border(
            bottom=Side(style="thin", color="17365D")
        )
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.border = header_border
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 28

        for row in sheet.iter_rows(min_row=2, max_col=5):
            row[0].number_format = "@"
            row[0].alignment = Alignment(horizontal="center", vertical="center")
            row[1].alignment = Alignment(horizontal="center", vertical="center")
            row[2].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            row[3].number_format = "dd/mm/yyyy hh:mm"
            row[3].alignment = Alignment(horizontal="center", vertical="center")
            row[4].alignment = Alignment(horizontal="center", vertical="center")

        widths = {"A": 12, "B": 24, "C": 58, "D": 22, "E": 18}
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width

        last_row = max(sheet.max_row, 2)
        table = Table(displayName="OrdenesTrabajoTable", ref=f"A1:E{last_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)

        status_range = f"E2:E{last_row}"
        sheet.conditional_formatting.add(
            status_range,
            FormulaRule(
                formula=['E2="Finalizada"'],
                fill=PatternFill("solid", fgColor="E2F0D9"),
                font=Font(color="375623"),
            ),
        )
        sheet.conditional_formatting.add(
            status_range,
            FormulaRule(
                formula=['E2="En progreso"'],
                fill=PatternFill("solid", fgColor="FFF2CC"),
                font=Font(color="7F6000"),
            ),
        )
        sheet.conditional_formatting.add(
            status_range,
            FormulaRule(
                formula=['E2="Cancelada"'],
                fill=PatternFill("solid", fgColor="FCE4D6"),
                font=Font(color="C00000"),
            ),
        )

        workbook.save(temp_excel)
        os.replace(temp_excel, EXCEL_PATH)
    except PermissionError as exc:
        raise ConsumanError(
            f"No se puede reemplazar {EXCEL_PATH.name} porque esta abierto en Excel. "
            "Cierre el archivo y vuelva a ejecutar. La extraccion JSON se completo."
        ) from exc
    except OSError as exc:
        raise ConsumanError(f"No se pudo generar el Excel: {exc}") from exc


def run() -> None:
    username, password = read_credentials()
    payload = build_payload()
    headless = os.getenv("CONSUMAN_HEADLESS", "0").strip().lower() in {
        "1", "true", "si", "yes"
    }

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=headless)
        context = browser.new_context()
        page = context.new_page()
        page.set_default_timeout(TIMEOUT_MS)
        try:
            login(page, username, password)
            data = extract_orders(page, payload)
            tasks = extract_tasks(page, payload)
            detail = build_task_records(data, tasks, task_columns(page))
            task_temp = TASK_OUTPUT_PATH.with_suffix(".json.tmp")
            task_temp.write_text(json.dumps(detail, ensure_ascii=False, indent=2), encoding="utf-8")
            os.replace(task_temp, TASK_OUTPUT_PATH)
            print(f"OK: {len(detail['records'])} tareas conservadas; campos de fecha: {detail['dateFields']}.")
            data = add_task_names(data, tasks)
            save_json_atomic(data)
            count = count_records(data)
            suffix = f" ({count} registros detectados)" if count is not None else ""
            print(f"OK: JSON guardado en {OUTPUT_PATH}{suffix}.")
            print("Generando Excel...")
            generate_excel(data)
            print(f"OK: Excel actualizado en {EXCEL_PATH}.")
        except PlaywrightTimeoutError as exc:
            raise ConsumanError(
                f"Consuman no respondio dentro de {TIMEOUT_MS // 1000} segundos. "
                f"URL actual: {page.url}"
            ) from exc
        finally:
            context.close()
            browser.close()


def main() -> int:
    try:
        run()
        return 0
    except ConsumanError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    except ValueError as exc:
        print(f"ERROR en el detalle de tareas: {exc}", file=sys.stderr)
        return 1
    except PlaywrightError as exc:
        print(f"ERROR de Playwright: {exc}", file=sys.stderr)
        return 1
    except OSError as exc:
        print(f"ERROR al guardar el archivo: {exc}", file=sys.stderr)
        return 1
    except KeyboardInterrupt:
        print("\nProceso cancelado por el usuario.", file=sys.stderr)
        return 130


if __name__ == "__main__":
    raise SystemExit(main())
